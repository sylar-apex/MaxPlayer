#!/usr/bin/env python3
"""Add every route, overlay, and setting row that the earlier sheets summarised."""
import re
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

PATH = "qa/MaxPlayer_UI_Control_Inventory.xlsx"
HEADER = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
FILL = PatternFill("solid", fgColor="9A3412")
ALT = PatternFill("solid", fgColor="FFF7ED")
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Border(
    left=Side(style="thin", color="D0D7DE"), right=Side(style="thin", color="D0D7DE"),
    top=Side(style="thin", color="D0D7DE"), bottom=Side(style="thin", color="D0D7DE"),
)

def style(ws, widths, fill):
    for cell in ws[1]:
        cell.font = HEADER
        cell.fill = fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for r in range(2, ws.max_row + 1):
        ws.row_dimensions[r].height = 36
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(r, c)
            cell.alignment = WRAP
            cell.border = THIN
            cell.font = Font(name="Calibri", size=10)
            if r % 2 == 0 and cell.fill.fgColor is None or (r % 2 == 0 and (cell.fill.fgColor is None or cell.fill.fill_type is None)):
                cell.fill = ALT
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.paperSize = ws.PAPERSIZE_TABLOID
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"
    ws.page_setup.horizontalDpi = 300
    ws.page_setup.verticalDpi = 300
    ws.oddHeader.left.text = ws.title
    ws.oddFooter.right.text = "Page &P of &N"
    ws.sheet_properties.tabColor = "9A3412"

# --- every page, from Routes.kt plus overlays that are not routes ---
pages = [
    ("Home", "Route", "Routes.Home", "App start", "Phone: Videos or Folders. Wide: Library or All videos.", "App screens", "Covered"),
    ("Folder", "Route", "Routes.Folder", "Home folder tap", "Title is the folder name. Empty: This folder is empty.", "App screens", "Covered"),
    ("Search", "Route", "Routes.Search", "Search icon, title-bar typing", "Search videos, recent searches, no-results line.", "App screens", "Covered"),
    ("Music", "Route", "Routes.Music", "Chip, sidebar", "Tabs Tracks, Playlists, Albums, Artists, Folders. Search and sort on Tracks.", "Controls + App screens", "Covered"),
    ("Music group", "In-page", "MusicScreen group state", "Album, artist, folder, Favourites, Recently Played, or a music playlist card", "Own header with Back, Play all, Shuffle, and that group's tracks. Not a separate route. Back or system back closes it.", "This sheet", "Added here"),
    ("Now Playing", "Overlay", "NowPlayingScreen", "Tap mini player", "Full music player. Not a navigation route.", "Controls", "Covered"),
    ("Mini player", "Overlay", "MiniPlayer", "isMusicSession", "Bar over Music only.", "Controls", "Covered"),
    ("Me / Profile", "Route", "Routes.Me", "Phone Account button", "Not in the sidebar. Wide users reach the same destinations from the sidebar.", "App screens", "Covered"),
    ("Theme", "Route", "Routes.Theme", "Me or sidebar", "App Theme modes and Classic Themes swatches.", "App screens", "Covered"),
    ("Settings index", "Route", "Routes.Settings", "Me, sidebar, player More", "List, Player, Decoder, Audio, Subtitle, General, Development, App Language.", "App screens", "Covered"),
    ("Settings page List", "Route", "Routes.SettingsPage List", "Settings index", "Appearance and Scan groups. Every row is on Every setting.", "Every setting", "Covered"),
    ("Settings page Player", "Route", "Routes.SettingsPage Player", "Settings index", "Theme, Style, Screen, Controls, Navigation plus playback settings.", "Every setting", "Covered"),
    ("Settings page Decoder", "Route", "Routes.SettingsPage Decoder", "Settings index", "HW/SW note and default decoder.", "Every setting", "Covered"),
    ("Settings page Audio", "Route", "Routes.SettingsPage Audio", "Settings index", "Audio player, volume, equalizer rows.", "Every setting", "Covered"),
    ("Settings page Subtitle", "Route", "Routes.SettingsPage Subtitle", "Settings index", "Text and Layout actions plus subtitle settings.", "Every setting", "Covered"),
    ("Settings page General", "Route", "Routes.SettingsPage General", "Settings index", "Links, edit, user-data actions.", "Every setting", "Covered"),
    ("Tuner Style", "Route", "Routes.Tuner 0", "Player Theme or Style, or in-player Display Settings", "Style preset and colours.", "Every setting", "Covered"),
    ("Tuner Screen", "Route", "Routes.Tuner 1", "Player Screen or Display Settings", "Orientation, fullscreen, brightness, OSD.", "Every setting", "Covered"),
    ("Tuner Controls", "Route", "Routes.Tuner 2", "Player Controls or Display Settings", "Touch, gestures, lock.", "Every setting", "Covered"),
    ("Tuner Navigation", "Route", "Routes.Tuner 3", "Player Navigation or Display Settings", "Seek speed, move buttons.", "Every setting", "Covered"),
    ("Tuner Text", "Route", "Routes.Tuner 4", "Subtitle Text or Display Settings", "Subtitle font and colour.", "Every setting", "Covered"),
    ("Tuner Layout", "Route", "Routes.Tuner 5", "Subtitle Layout or Display Settings", "Alignment and margins.", "Every setting", "Covered"),
    ("Video playlists", "Route", "Routes.Playlists", "Chip, sidebar, Me", "Title VIDEO PLAYLISTS.", "App screens", "Covered"),
    ("Playlist detail", "Route", "Routes.PlaylistDetail", "Playlist row", "Play all, rename, delete, remove item.", "App screens", "Covered"),
    ("Private folder", "Route", "Routes.PrivateFolder", "Privacy chip, sidebar, Me", "PIN gate then file list.", "App screens", "Covered"),
    ("Media manager", "Route", "Routes.MediaManager", "Me or sidebar Storage", "Storage cards and list shortcuts.", "App screens", "Covered"),
    ("Recently Played list", "Route", "Routes.MediaList Recent", "Media manager card", "File list. Empty: Nothing here.", "App screens", "Covered"),
    ("Large Files list", "Route", "Routes.MediaList Large", "Media manager card", "File list.", "App screens", "Covered"),
    ("Recently Added list", "Route", "Routes.MediaList Added", "Media manager card", "File list.", "App screens", "Covered"),
    ("Haven't Played list", "Route", "Routes.MediaList Unplayed", "Haven't Played row", "File list.", "App screens", "Covered"),
    ("Help", "Route", "Routes.Help", "Me, sidebar, player FAQ", "Five FAQ topics.", "App screens", "Covered"),
    ("Local network", "Route", "Routes.LocalNetwork", "Me, or sidebar item named Network stream", "Title Networks. Plus toasts not supported.", "App screens", "Covered"),
    ("Unavailable File Transfer", "Route", "Routes.Unavailable", "Me tile. Home onFileTransfer is unused by the current chips.", "Reason: MXShare is not implemented.", "App screens", "Covered"),
    ("Unavailable Cloud Drive", "Route", "Routes.Unavailable", "Me tile", "OAuth not implemented.", "App screens", "Covered"),
    ("Unavailable WhatsApp Status Saver", "Route", "Routes.Unavailable", "Me row", "SAF permission not implemented.", "App screens", "Covered"),
    ("Unavailable Legal", "Route", "Routes.Unavailable", "Me row", "Info text only.", "App screens", "Covered"),
    ("Unavailable Development", "Route", "Routes.Unavailable", "Settings index", "Diagnostics not exposed.", "App screens", "Covered"),
    ("Video player (Android)", "Activity", "PlayerActivity", "Play a video on Android", "Separate activity. Same VideoPlayerScreen. Also opens external video and stream intents.", "Controls", "Covered"),
    ("Video player (embedded)", "Overlay", "App.EmbeddedPlayer", "Desktop, web, iOS", "Covers the library. Close returns to the page underneath.", "Controls", "Covered"),
    ("Initialising player", "Overlay", "PlayerInitialising", "Embedded player before the engine is ready", "Black screen. Text Initialising player… and Loading video codecs for the first time. Clicks do not reach the library.", "This sheet", "Added here"),
    ("Android permission prompt", "System dialog", "MainActivity", "First launch on Android if media or notification permission is missing", "System permission dialog, not an in-app page. API 33+: videos, audio, images, notifications. Older: storage. Deny still opens Home; scan stays empty until granted.", "This sheet", "Added here"),
    ("Delete consent", "System dialog", "requestConsent", "Android delete when MediaStore requires consent", "System dialog. Allow completes the delete. Deny leaves the file.", "This sheet", "Added here"),
    ("Playback notification", "System UI", "PlaybackService", "Android background play", "Media3 notification and lock-screen controls. No custom labels in app code. Tap opens music activity if a music session, else the video activity. Hidden from buttons if media buttons setting is off. Task removed stops the service if not playing.", "This sheet", "Added here"),
    ("View menu", "Dialog", "ViewMenuDialog", "View Menu or View options", "Not a route. Sort, fields, advanced filters.", "App screens", "Covered"),
    ("File actions sheet", "Dialog", "FileActionsSheet", "Row more menu", "Not a route.", "App screens", "Covered"),
    ("Network Stream dialog", "Dialog", "NetworkStreamDialog", "Me tile, or player overflow", "Not the sidebar Network stream item.", "Controls", "Covered"),
    ("Title bar", "Chrome", "TitleBar", "Wide windows only", "Back, Forward, Refresh, Search videos. Phones do not have this bar.", "App screens", "Covered"),
    ("Sidebar", "Chrome", "Sidebar", "Wide windows only", "LIBRARY and MORE. Phones use chips and Me.", "App screens", "Covered"),
]

# settings parser
def parse_settings(path, registry):
    text = open(path).read()
    out = []
    # BoolSetting("key", "Title", default, "summary"?)
    for m in re.finditer(r'(BoolSetting|ChoiceSetting|IntSetting|FloatSetting|TextSetting|ColorSetting)\((.*?)\)\s*(?:,|\))', text, re.S):
        kind, body = m.group(1), m.group(2)
        strings = re.findall(r'"((?:\\.|[^"\\])*)"', body)
        if len(strings) < 2:
            continue
        key, title = strings[0], strings[1]
        summary = ""
        choices = ""
        default = ""
        if kind == "BoolSetting":
            dm = re.search(r',\s*(true|false)', body)
            default = dm.group(1) if dm else ""
            if len(strings) > 2:
                summary = strings[2]
        elif kind == "ChoiceSetting":
            default = strings[2] if len(strings) > 2 else ""
            pairs = re.findall(r'"([^"]+)"\s+to\s+"([^"]+)"', body)
            choices = "; ".join(f"{a} = {b}" for a, b in pairs)
            if len(strings) > 2 + len(pairs) * 2:
                summary = strings[-1]
        elif kind == "IntSetting":
            dm = re.search(r',\s*(-?\d+)', body)
            default = dm.group(1) if dm else ""
            rng = re.search(r'(-?\d+)\.\.(-?\d+)', body)
            unit = strings[2] if len(strings) > 2 and strings[2] not in (summary,) else ""
            # unit is last quoted if present after range
            qs = strings[2:]
            if qs:
                # first extra may be summary, last short may be unit
                if len(qs) == 1 and len(qs[0]) < 8 and " " not in qs[0]:
                    unit = qs[0]
                else:
                    summary = qs[0]
                    if len(qs) > 1:
                        unit = qs[-1]
            if rng:
                summary = (summary + f" Range {rng.group(1)}..{rng.group(2)}." + (f" Unit {unit}." if unit else "")).strip()
        elif kind == "TextSetting":
            default = strings[2] if len(strings) > 2 else ""
            if len(strings) > 3:
                summary = strings[3]
        elif kind == "ColorSetting":
            cm = re.search(r'(0x[0-9A-Fa-f]+)', body)
            default = cm.group(1) if cm else ""
        elif kind == "FloatSetting":
            dm = re.search(r',\s*(-?\d+(?:\.\d+)?)', body)
            default = dm.group(1) if dm else ""
            if len(strings) > 2:
                summary = strings[2]
        out.append((registry, kind.replace("Setting", ""), key, title, default, summary, choices))
    return out

settings = parse_settings("kmediaplayer/core/src/commonMain/kotlin/org/app/kmediaplayer/settings/PlayerKeys.kt", "PlayerKeys")
settings += parse_settings("appcore/sharedlogic/src/commonMain/kotlin/org/app/maxplayerpro/sharedlogic/settings/AppKeys.kt", "AppKeys")

# where each key is shown
where = {
    "list.view_options": "Saved by View Menu. Not shown as its own row.",
    "app_theme": "App Theme page, not a settings row.",
    "accent_color": "App Theme Classic Themes, and PlayerKeys. Not a settings row.",
}
# tuner membership from Tuner.kt is enough to say Display Settings + Settings player links

wb = load_workbook(PATH)
for name in ("Every page", "Every setting"):
    if name in wb.sheetnames:
        del wb[name]

ep = wb.create_sheet("Every page", 1)
ep.append(["Page", "Kind", "Route / type", "How to open", "What is on it", "Already listed in", "Status"])
for row in pages:
    ep.append(row)
style(ep, [32, 16, 28, 42, 62, 28, 14], FILL)
dv = DataValidation(type="list", formula1='"Pass,Fail,Blocked,Not run"', allow_blank=True)
ep.cell(1, 8, "QA result").font = HEADER
ep.cell(1, 8).fill = FILL
ep.column_dimensions["H"].width = 14
dv.add(f"H2:H{ep.max_row}")
ep.add_data_validation(dv)
ep.auto_filter.ref = f"A1:H{ep.max_row}"
ep.conditional_formatting.add(f"H2:H{ep.max_row}", FormulaRule(formula=['$H2="Fail"'], fill=PatternFill("solid", fgColor="FECACA")))

es = wb.create_sheet("Every setting", 2)
es.append(["Registry", "Control type", "Storage key", "Exact title", "Default", "Summary / range", "Choices", "How it behaves", "QA result"])
for reg, kind, key, title, default, summary, choices in settings:
    if kind == "Bool":
        how = "Settings page: row tap or checkbox toggles it. In-player Display Settings uses a checkbox."
    elif kind == "Choice":
        how = "Opens a radio dialog titled with the setting name. Picking a choice saves and closes. Cancel closes without saving. Empty label shows Default."
    elif kind == "Int":
        how = "Slider if the range is 400 or less; otherwise a number dialog OK / CANCEL."
    elif kind == "Float":
        how = "Number dialog OK / CANCEL."
    elif kind == "Text":
        how = "Text dialog OK / CANCEL. Folders (list.scan_folders) is multiline. Empty value shows the summary instead."
    elif kind == "Color":
        how = "Opens a colour dialog. Swatches, Opacity slider, OK / Cancel."
    else:
        how = ""
    note = where.get(key, "")
    if note:
        how = note + " " + how
    es.append((reg, kind, key, title, default, summary, choices, how, ""))
style(es, [14, 14, 36, 42, 16, 48, 42, 55, 14], PatternFill("solid", fgColor="1E3A8A"))
for cell in es[1]:
    cell.fill = PatternFill("solid", fgColor="1E3A8A")
dv2 = DataValidation(type="list", formula1='"Pass,Fail,Blocked,Not run"', allow_blank=True)
dv2.add(f"I2:I{es.max_row}")
es.add_data_validation(dv2)
es.auto_filter.ref = f"A1:I{es.max_row}"

# setting control mechanics as pages
cover = wb["Read me"]
cover["A19"] = "Complete route list"
cover["A19"].font = Font(bold=True, color="9A3412")
cover["A20"] = "Every page lists all 17 navigation routes, the six tuner tabs, the four media lists, and the overlays that are not routes (player, notification, permission). Every setting lists every PlayerKeys and AppKeys row."
cover["A20"].alignment = WRAP
cover.merge_cells("A20:F20")
cover.row_dimensions[20].height = 36
cover["A21"] = "Pages"
cover["B21"] = f"=COUNTA('Every page'!A:A)-1"
cover["A22"] = "Settings"
cover["B22"] = f"=COUNTA('Every setting'!A:A)-1"
cover["A21"].font = Font(bold=True)
cover["A22"].font = Font(bold=True)

wb.save(PATH)
print("pages", len(pages), "settings", len(settings))
